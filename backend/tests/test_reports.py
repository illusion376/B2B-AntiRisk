import codecs
import json
import uuid
from datetime import datetime

import pymupdf
import pytest
from docx import Document as DocxDocument

from app.models import RiskRule
from app.services import reports
from app.services.extraction import extract_pages
from app.services.quotes import locate_quote
from app.services.reports.data import ReportData, ReportDocument, ReportFinding
from app.services.reports.docx_builder import BUILDERS


def _data(contract_pdf) -> ReportData:
    pages = [{"page_number": p.page_number, "words": p.words} for p in extract_pages(contract_pdf)]
    match = locate_quote(pages, "штраф в размере 0,1% от стоимости Контракта", (1, 1))

    def finding(number, severity, title, **kw):
        base = dict(number=number, severity=severity, category="Штрафы", title=title, short_description="Суть",
                    page_number=1, clause="6.2", exact_quote=match.text, comment="Обоснование",
                    counter_proposal="Новая редакция", legal_reference="ст. 34 44-ФЗ", review_status="NEW",
                    reviewer_comment=None, quote_verified=True, source="LLM", highlights=match.highlights)
        base.update(kw)
        return ReportFinding(**base)

    doc = ReportDocument(
        id=uuid.uuid4(), file_name="Проект контракта.pdf", relative_path="Проект контракта.pdf", status="COMPLETED",
        total_pages=2, is_scanned=False, law_type="44-FZ", risk_score=37, light="RED",
        preview_path=str(contract_pdf), error_message=None,
        findings=[
            finding(1, "RED", "Неограниченная ответственность"),
            finding(2, "YELLOW", "Срок поставки", exact_quote=None, highlights=[], page_number=2, clause="7.2"),
            finding(3, "LOW", "Подсудность", exact_quote=None, highlights=[], page_number=2, clause="11"),
            finding(None, "GREEN", "Срок оплаты", exact_quote=None, highlights=[], page_number=None, clause=None,
                    comment="Срок оплаты 7 рабочих дней"),
        ],
    )
    rule = RiskRule(id="r1", title="Неустойка", category="Штрафы", severity="RED", legal_reference="ст. 34")
    return ReportData(
        analysis_id=uuid.uuid4(), title="Проект контракта", original_filename="Проект контракта.pdf",
        created_at=datetime.now(), generated_at=datetime.now(), author="Кирилл Иванов", company="ООО Поставщик",
        rules_checked=3, documents=[doc], rules=[rule],
    )


@pytest.mark.parametrize("mode", ["brief", "detailed", "protocol"])
def test_docx_modes(contract_pdf, tmp_path, mode):
    data = _data(contract_pdf)
    path = BUILDERS[mode](data, tmp_path / f"{mode}.docx")
    docx = DocxDocument(path)
    text = "\n".join(p.text for p in docx.paragraphs)
    cells = "\n".join(c.text for t in docx.tables for r in t.rows for c in r.cells)
    if mode != "protocol":
        assert "Неограниченная ответственность" in cells + text
        assert "п. 6.2" in cells + text and "Низкий риск" in cells + text
    else:
        assert "ПРОТОКОЛ РАЗНОГЛАСИЙ" in text
        assert "Новая редакция" in cells and "штраф в размере 0,1%" in cells
    if mode == "detailed":
        assert "Срок оплаты 7 рабочих дней" in text  # пройденные проверки


def test_json_and_annotated(contract_pdf):
    data = _data(contract_pdf)
    report = reports.generate(data, "brief", "json")
    payload = json.loads(report.path.read_text(encoding="utf-8"))
    assert payload["traffic_light"] == "RED" and len(payload["documents"][0]["findings"]) == 3

    table = reports.generate(data, "brief", "csv").path.read_bytes()
    assert table.startswith(codecs.BOM_UTF8)  # BOM для Excel
    rows = table.decode("utf-8-sig").splitlines()
    assert any('"Критические замечания";"Неограниченная ответственность"' in r for r in rows)
    assert not any("Срок оплаты" in r for r in rows)  # краткий режим — без пройденных проверок

    annotated = reports.generate(data, "annotated", "pdf")
    with pymupdf.open(annotated.path) as pdf:
        page = pdf[0]
        types = [a.type[1] for a in page.annots()]
        assert "Highlight" in types and "Text" in types

    with pytest.raises(reports.ReportError):
        reports.generate(data, "annotated", "docx")


def test_csv_formula_injection_is_neutralized():
    from app.services.reports import csv_cell

    assert csv_cell("=HYPERLINK(\"http://x\")") == "'=HYPERLINK(\"http://x\")"
    assert csv_cell("  +7 (999)") == "'  +7 (999)"
    assert csv_cell("@SUM(A1)") == "'@SUM(A1)"
    assert csv_cell("Срок оплаты — 30 дней") == "Срок оплаты — 30 дней"
    assert csv_cell(5) == 5
