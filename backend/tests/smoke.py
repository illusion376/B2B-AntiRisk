"""Сквозная проверка запущенной системы (не юнит-тест, pytest его не собирает).

Запуск внутри контейнера после `docker-compose up --build`:
    docker compose exec backend python -m tests.smoke

Скрипт сам собирает ZIP с тестовой документацией (текстовый PDF с рискованными пунктами,
страница-скан для OCR, DOCX для конвертации), создаёт проект, загружает архив, ждёт обработки
и проверяет все основные эндпоинты в словаре фронтенда. Отчёты сохраняются в /tmp/smoke_reports.
"""
import io
import os
import sys
import time
import zipfile
from pathlib import Path

import httpx
import pymupdf
from docx import Document as DocxDocument

API = os.environ.get("SMOKE_API", "http://localhost:8000")
OUT = Path(os.environ.get("SMOKE_OUT", "/tmp/smoke_reports"))
FONT = next((f for f in (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "C:/Windows/Fonts/arial.ttf",
) if Path(f).exists()), None)

CONTRACT = [
    """ПРОЕКТ КОНТРАКТА № 0123 на поставку офисной мебели
Контракт заключается в соответствии с Федеральным законом от 05.04.2013 № 44-ФЗ.

1. ПРЕДМЕТ КОНТРАКТА
1.1. Поставщик обязуется поставить офисную мебель в соответствии со Спецификацией, а Заказчик обязуется принять и оплатить Товар.

4. ПОРЯДОК ОПЛАТЫ
4.3. Оплата производится в течение 30 календарных дней с даты подписания Заказчиком документа о приемке.

5. ПРИЕМКА ТОВАРА
5.2. Заказчик вправе отказаться от подписания документа о приемке без указания причин.""",
    """6. ОТВЕТСТВЕННОСТЬ СТОРОН
6.1. Стороны несут ответственность за неисполнение обязательств в соответствии с законодательством Российской Федерации.
6.2. Поставщик уплачивает Заказчику штраф в размере 0,1% от цены Контракта за каждый день просрочки исполнения обязательств, при этом общая сумма неустойки не ограничена ценой Контракта.
6.3. Поставщик обязан устранить недостатки Товара в течение 1 (одного) дня с момента устного уведомления Заказчика.

8. РАСТОРЖЕНИЕ КОНТРАКТА
8.1. Заказчик вправе в одностороннем порядке отказаться от исполнения Контракта по любому основанию, в том числе при отсутствии нарушений со стороны Поставщика.""",
]
SCANNED_PAGE = """9. ОБЕСПЕЧЕНИЕ ИСПОЛНЕНИЯ КОНТРАКТА
9.1. Размер обеспечения исполнения Контракта составляет 30% от начальной (максимальной) цены контракта.
9.2. Обеспечение предоставляется в виде денежных средств на счет Заказчика."""

failures: list[str] = []


def check(condition: bool, message: str) -> None:
    print(("  OK   " if condition else "  FAIL ") + message)
    if not condition:
        failures.append(message)


def build_contract_pdf() -> bytes:
    if FONT is None:
        sys.exit("Не найден шрифт с кириллицей для генерации тестового PDF")
    doc = pymupdf.open()
    for text in CONTRACT:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(60, 60, 540, 780), text, fontsize=11, fontname="cyr", fontfile=FONT)

    # Страница-скан: рендерим текст в картинку и вставляем только изображение (без текстового слоя)
    tmp = pymupdf.open()
    tmp_page = tmp.new_page()
    tmp_page.insert_textbox(pymupdf.Rect(60, 60, 540, 780), SCANNED_PAGE, fontsize=12, fontname="cyr", fontfile=FONT)
    image = tmp_page.get_pixmap(dpi=200, colorspace=pymupdf.csGRAY).tobytes("png")
    scan = doc.new_page()
    scan.insert_image(scan.rect, stream=image)
    return doc.tobytes()


def build_spec_docx() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Техническое задание", level=1)
    doc.add_paragraph("1. ТРЕБОВАНИЯ К ТОВАРУ")
    doc.add_paragraph("1.1. Стол офисный производства IKEA, модель BEKANT. Поставка эквивалента не допускается.")
    doc.add_paragraph("1.2. Качество товара должно быть высоким, дизайн — современным.")
    doc.add_paragraph("2. УСЛОВИЯ ПОСТАВКИ")
    doc.add_paragraph("2.1. Доставка, разгрузка, подъём на этаж и сборка мебели по 15 адресам Заказчика "
                      "осуществляются силами и за счёт Поставщика.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("Документация/Проект контракта.pdf", build_contract_pdf())
        zf.writestr("Документация/Техническое задание.docx", build_spec_docx())
        zf.writestr("Документация/readme.exe", b"not a document")
    return buf.getvalue()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    client = httpx.Client(base_url=API, timeout=120)

    print("1. Health, пользователь, правила")
    health = client.get("/api/health").json()
    print("    ", health)
    check(health["database"] and health["queue"], "БД и очередь доступны")
    me = client.get("/api/users/me").json()
    check(bool(me.get("initials")), f"текущий пользователь: {me.get('full_name')}")
    rules = client.get("/api/rules", params={"enabled": True}).json()
    check(len(rules) >= 20, f"включённых правил: {len(rules)}")
    check(all(r["severity"] in ("critical", "warning", "low") for r in rules), "уровни правил в словаре фронтенда")

    print("2. Проект и загрузка файлов")
    title = f"Smoke-тест {time.strftime('%H:%M:%S')}"
    project = client.post("/api/projects", json={"title": title, "description": "Сквозная проверка"}).json()
    pid = project["id"]
    duplicate = client.post("/api/projects", json={"title": title.upper()})
    check(duplicate.status_code == 409, "повторное название проекта отклоняется (без учёта регистра)")

    started = time.monotonic()
    response = client.post(f"/api/projects/{pid}/files", files=[
        ("files", ("Документация.zip", build_zip(), "application/zip")),
        ("files", ("пустой.pdf", b"", "application/pdf")),
    ])
    check(response.status_code == 202, f"загрузка файлов в проект -> {response.status_code}")
    upload = response.json()
    check(len(upload["files"]) == 1 and len(upload["errors"]) == 1,
          f"ZIP принят, пустой файл отклонён: {[e['detail'] for e in upload['errors']]}")
    file_id = upload["files"][0]["id"]
    docs = {d["relative_path"]: d for d in upload["files"][0]["documents"]}
    print("     файлы архива:", {k: v["label"] for k, v in docs.items()})
    check(len(docs) == 3, "в ответе сразу видны все 3 файла архива")
    check(docs.get("Документация/readme.exe", {}).get("phase") == "unsupported", "exe помечен как неподдерживаемый")

    check(upload["files"][0]["phase"] == "uploaded", "загрузка не запускает анализ")
    response = client.post(f"/api/projects/{pid}/start")
    check(response.status_code == 202 and response.json()["documents"] == 2, "явный запуск двух поддерживаемых документов")

    print("3. Ожидание обработки")
    last = None
    while True:
        project = client.get(f"/api/projects/{pid}").json()
        file = project["files"][0]
        if file["label"] != last:
            print(f"     {time.monotonic() - started:6.1f}s  {file['label']}")
            last = file["label"]
        if file["phase"] in ("ready", "failed"):
            break
        if time.monotonic() - started > 600:
            check(False, "обработка уложилась в 10 минут")
            break
        time.sleep(2)
    elapsed = time.monotonic() - started
    check(file["phase"] == "ready", f"файл обработан за {elapsed:.0f} с, светофор={file['traffic_light']}, "
                                    f"риск={file['risk_score']}, замечания={file['counts']}")
    check(project["processing_count"] == 0, "в проекте нет файлов в обработке")
    for d in file["documents"]:
        print(f"     {d['relative_path']}: {d['label']}, стр. {d['total_pages']}, OCR-стр. {d['ocr_pages']}, "
              f"закон {d['law_type']}, {d['counts']}")
        if d["phase"] != "unsupported":
            check(d["phase"] == "ready", f"{d['file_name']} обработан")

    contract = next(d for d in file["documents"] if d["file_name"].endswith(".pdf"))
    spec = next(d for d in file["documents"] if d["file_name"].endswith(".docx"))
    cid = contract["id"]

    print("4. Документ: OCR, текст страниц, навигация, поиск")
    check(contract["ocr_pages"] >= 1, f"скан распознан OCR (уверенность {contract['ocr_confidence']})")
    check(contract["law_type"] == "44-FZ", "закон определён как 44-ФЗ")
    page2 = client.get(f"/api/documents/{cid}/pages/2").json()
    print("     стр. 2:", [(s["title"], [p["clause"] for p in s["paragraphs"]]) for s in page2["sections"]])
    check(bool(page2["sections"]) and page2["sections"][0]["title"].startswith("6."), "текст страницы разделами и пунктами")
    check(any(p["finding_ids"] for s in page2["sections"] for p in s["paragraphs"]), "замечания привязаны к абзацам")
    hits = client.get(f"/api/documents/{cid}/search", params={"q": "обеспечения исполнения"}).json()
    check(any(h["page"] == 3 for h in hits["hits"]), "поиск находит текст на странице-скане")
    check(any(h["clause"] for h in hits["hits"]), "в результатах поиска есть номер пункта")
    preview = client.get(f"/api/documents/{cid}/file")
    check(preview.status_code == 200 and preview.content.startswith(b"%PDF"), "PDF документа отдаётся")
    spec_preview = client.get(f"/api/documents/{spec['id']}/file")
    check(spec_preview.status_code == 200 and spec_preview.content.startswith(b"%PDF"), "DOCX сконвертирован в PDF")
    thumb = client.get(f"/api/documents/{cid}/pages/1/thumbnail", params={"width": 160})
    check(thumb.status_code == 200 and thumb.content[:4] == b"\x89PNG", "миниатюра страницы")
    outline = client.get(f"/api/documents/{cid}/outline").json()
    print("     навигация:", [s["title"] for s in outline])
    check(len(outline) >= 4, "в навигации найдены разделы договора")

    print("5. Замечания")
    findings = client.get(f"/api/documents/{cid}/findings").json()
    for group in findings["groups"]:
        print(f"     {group['label']} ({group['count']})")
        for f in group["items"]:
            if f["severity"] == "ok":
                continue
            print(f"       №{f['number']} [{f['source']}] {f['title']} — п. {f['clause']}, стр. {f['page']}, "
                  f"цитата {'✓' if f['quote_verified'] else '✗'}: {f['quote'][:60]}")
    issues = [f for g in findings["groups"] if g["severity"] != "ok" for f in g["items"]]
    check(len(issues) > 0, f"найдено замечаний: {len(issues)}")
    check(all(f["page"] for f in issues), "у каждого замечания есть страница")
    verified = [f for f in issues if f["quote_verified"]]
    check(len(verified) >= len(issues) * 0.7, f"цитаты подтверждены в тексте: {len(verified)} из {len(issues)}")
    check(all(f["highlights"] for f in verified), "у подтверждённых цитат есть координаты подсветки")
    if health["llm"].startswith("heuristic"):
        print("     (LLM не настроена — замечания найдены эвристикой; для оценки качества задайте LLM_BASE_URL)")

    if issues:
        patched = client.patch(f"/api/findings/{issues[0]['id']}", json={"status": "accepted"})
        check(patched.status_code == 200 and patched.json()["status"] == "accepted", "смена статуса проверки")
        filtered = client.get(f"/api/documents/{cid}/findings", params={"status": "accepted"}).json()
        check(filtered["total"] == 1 and filtered["statuses"]["accepted"] == 1, "фильтр по статусу")

    print("6. Правило из интерфейса")
    rule = client.post("/api/rules", json={"title": "Аванс", "description": "Аванс не предусмотрен",
                                           "category": "Оплата", "severity": "low"}).json()
    check(rule.get("severity") == "low" and bool(rule.get("llm_prompt")), "правило создано из полей формы фронтенда")
    toggled = client.patch(f"/api/rules/{rule['id']}", json={"enabled": False}).json()
    check(toggled.get("enabled") is False, "правило отключено")
    check(client.delete(f"/api/rules/{rule['id']}").status_code == 204, "правило удалено")

    print("7. Отчёты")
    for mode, fmt in [("brief", "docx"), ("detailed", "docx"), ("protocol", "docx"), ("brief", "pdf"),
                      ("brief", "csv"), ("detailed", "json"), ("annotated", "pdf")]:
        report = client.get(f"/api/documents/{cid}/report", params={"mode": mode, "format": fmt})
        ok = report.status_code == 200 and len(report.content) > 300
        path = OUT / f"{mode}.{fmt}"
        if ok:
            path.write_bytes(report.content)
        check(ok, f"отчёт {mode}.{fmt} ({len(report.content)} байт) -> {path}")

    print("8. История")
    history = client.get("/api/history", params={"limit": 30}).json()
    for entry in history[:6]:
        print(f"     {entry['title']}: {entry['detail']}")
    actions = {e["action"] for e in history}
    check({"PROJECT_CREATED", "ANALYSIS_CREATED", "ANALYSIS_COMPLETED", "FINDING_REVIEWED"} <= actions,
          "в истории есть создание проекта, загрузка, завершение проверки и смена статуса")

    print()
    if failures:
        print(f"ПРОВАЛЕНО проверок: {len(failures)}")
        for f in failures:
            print("  -", f)
        sys.exit(1)
    print(f"Все проверки пройдены. Проект «{title}», id={pid}, файл id={file_id}")


if __name__ == "__main__":
    main()
