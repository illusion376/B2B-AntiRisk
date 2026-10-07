"""Отчёты в DOCX: краткий, подробный и протокол разногласий."""
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.services.reports.data import ReportData, ReportDocument, ReportFinding
from app.vocab import REVIEW_LABELS, SEVERITY_SHORT_LABELS

LIGHT_LABEL = {"RED": "Красный — критичные риски", "YELLOW": "Жёлтый — требует внимания", "GREEN": "Зелёный — без рисков"}
SEVERITY_LABEL = SEVERITY_SHORT_LABELS
SEVERITY_FILL = {"RED": "F8D7D5", "YELLOW": "FDEFC8", "LOW": "E8F5E9", "GREEN": "DDF0DE"}
SEVERITY_TEXT = {"RED": RGBColor(0xB7, 0x1C, 0x1C), "YELLOW": RGBColor(0x9A, 0x67, 0x00),
                 "LOW": RGBColor(0x2E, 0x7D, 0x32), "GREEN": RGBColor(0x2E, 0x7D, 0x32)}
REVIEW_LABEL = REVIEW_LABELS
LAW_LABEL = {"44-FZ": "44-ФЗ", "223-FZ": "223-ФЗ"}
VERDICT = {
    "RED": "Выявлены критичные условия. Рекомендуется направить запрос разъяснений или подготовить "
           "протокол разногласий до подачи заявки / подписания контракта.",
    "YELLOW": "Критичных условий не выявлено, но есть пункты, требующие внимания при расчёте цены и планировании исполнения.",
    "GREEN": "Существенных рисков по проверенным правилам не выявлено.",
}
DISCLAIMER = ("Отчёт сформирован автоматически ИИ-ассистентом «Светофор рисков» и не является юридическим "
              "заключением. Замечания необходимо проверить специалисту.")


# ---------- примитивы ----------

def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _new_document(landscape: bool = False) -> DocxDocument:
    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(11)
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2))
    section.top_margin = section.bottom_margin = Cm(1.8)
    return doc


def _cell_text(cell, text: str, bold: bool = False, size: int | None = None, color: RGBColor | None = None,
               italic: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text or "—")
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def _table(doc, headers: list[str], widths: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        _cell_text(cell, header, bold=True, size=10)
        _shade(cell, "EDEDED")
    _repeat_header(table.rows[0])
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return table


def _add_row(table, values: list[str], widths: list[float], size: int = 10):
    row = table.add_row()
    for i, value in enumerate(values):
        _cell_text(row.cells[i], value, size=size)
        row.cells[i].width = Cm(widths[i])
    return row


def clause_label(clause: str | None) -> str | None:
    """«6.2» -> «п. 6.2»; ненумерованные значения — как есть."""
    if not clause:
        return None
    return f"п. {clause}" if clause[0].isdigit() else clause


def _location(f: ReportFinding) -> str:
    parts = [p for p in (clause_label(f.clause), f"стр. {f.page_number}" if f.page_number else None) if p]
    return ", ".join(parts) or "—"


def _header(doc, data: ReportData, title: str) -> None:
    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    rows = [
        ("Пакет документов", data.title),
        ("Исходный файл", data.original_filename),
        ("Дата загрузки", data.created_at.strftime("%d.%m.%Y %H:%M") if data.created_at else "—"),
        ("Дата отчёта", data.generated_at.strftime("%d.%m.%Y %H:%M")),
        ("Проверено правил", str(data.rules_checked)),
        ("Подготовил", ", ".join(p for p in (data.author, data.company) if p) or "—"),
    ]
    for label, value in rows:
        cells = meta.add_row().cells
        _cell_text(cells[0], label, bold=True)
        _cell_text(cells[1], value)
        cells[0].width, cells[1].width = Cm(5), Cm(12)
    doc.add_paragraph()


def _summary(doc, data: ReportData) -> None:
    doc.add_heading("Итог проверки", level=2)
    findings = data.all_findings
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    light = data.light
    values = [
        (f"Светофор: {LIGHT_LABEL[light]}", SEVERITY_FILL[light]),
        (f"Критичных: {sum(f.severity == 'RED' for f in findings)}", SEVERITY_FILL["RED"]),
        (f"Требуют внимания: {sum(f.severity == 'YELLOW' for f in findings)}"
         + (f", низкий риск: {low}" if (low := sum(f.severity == 'LOW' for f in findings)) else ""),
         SEVERITY_FILL["YELLOW"]),
        (f"Индекс риска: {data.risk_score} из 100", "EDEDED"),
    ]
    for cell, (text, fill) in zip(table.rows[0].cells, values):
        _cell_text(cell, text, bold=True)
        _shade(cell, fill)
    doc.add_paragraph(VERDICT[light])


def _document_title(doc, rd: ReportDocument, multi: bool) -> None:
    if not multi:
        return
    doc.add_heading(rd.relative_path or rd.file_name, level=2)
    info = f"Страниц: {rd.total_pages}"
    if rd.law_type:
        info += f" · {LAW_LABEL.get(rd.law_type, rd.law_type)}"
    if rd.is_scanned:
        info += " · распознан OCR"
    if rd.status != "COMPLETED":
        info += f" · не обработан: {rd.error_message or rd.status}"
    doc.add_paragraph(info)


def _footer(doc) -> None:
    doc.add_paragraph()
    p = doc.add_paragraph(DISCLAIMER)
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)


# ---------- режимы ----------

def build_brief(data: ReportData, target: Path) -> Path:
    doc = _new_document()
    _header(doc, data, "Краткий отчёт о проверке документации")
    _summary(doc, data)
    multi = len(data.documents) > 1
    widths = [1.0, 2.6, 8.4, 3.0, 1.5]
    doc.add_heading("Замечания", level=2)
    for rd in data.documents:
        _document_title(doc, rd, multi)
        issues = rd.issues
        if not issues:
            doc.add_paragraph("Замечаний нет." if rd.status == "COMPLETED" else "Документ не проверен.")
            continue
        table = _table(doc, ["№", "Уровень", "Замечание", "Пункт", "Стр."], widths)
        for f in issues:
            row = _add_row(table, [
                str(f.number or ""), SEVERITY_LABEL[f.severity],
                f"{f.title}. {f.short_description or ''}".strip(), clause_label(f.clause) or "—", str(f.page_number or "—"),
            ], widths)
            _shade(row.cells[1], SEVERITY_FILL[f.severity])
    _footer(doc)
    doc.save(target)
    return target


def build_detailed(data: ReportData, target: Path) -> Path:
    doc = _new_document()
    _header(doc, data, "Подробный отчёт о проверке документации")
    _summary(doc, data)
    multi = len(data.documents) > 1

    for rd in data.documents:
        if multi:
            _document_title(doc, rd, multi)
        for f in rd.issues:
            p = doc.add_paragraph()
            run = p.add_run(f"№{f.number}. {SEVERITY_LABEL[f.severity]}: {f.title}")
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = SEVERITY_TEXT[f.severity]
            meta = [_location(f)]
            if f.category:
                meta.append(f.category)
            meta.append(f"статус: {REVIEW_LABEL.get(f.review_status, f.review_status)}")
            doc.add_paragraph(" · ".join(meta)).runs[0].font.size = Pt(9)

            if f.exact_quote:
                quote_table = doc.add_table(rows=1, cols=1)
                quote_table.style = "Table Grid"
                cell = quote_table.rows[0].cells[0]
                _cell_text(cell, f"«{f.exact_quote}»", italic=True, size=10)
                _shade(cell, SEVERITY_FILL[f.severity])
                if not f.quote_verified:
                    doc.add_paragraph("Цитата не найдена в тексте дословно — сверьте с документом.").runs[0].italic = True
            _labeled(doc, "Почему это риск", f.comment)
            if f.counter_proposal:
                _labeled(doc, "Рекомендация", f.counter_proposal)
            if f.legal_reference:
                _labeled(doc, "Норма", f.legal_reference)
            if f.reviewer_comment:
                _labeled(doc, "Комментарий проверяющего", f.reviewer_comment)
            if f.source == "HEURISTIC":
                _labeled(doc, "Примечание", "Найдено без LLM (по ключевым словам) — требуется ручная проверка")
        if rd.status == "COMPLETED" and not rd.issues:
            doc.add_paragraph("Замечаний нет.")

        if rd.green:
            doc.add_heading("Проверено, нарушений не выявлено", level=3)
            for f in rd.green:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"{f.title}. ").bold = True
                p.add_run(f.comment)

    if data.rules:
        doc.add_heading("Приложение. Правила проверки", level=2)
        widths = [6.5, 3.5, 2.2, 4.8]
        table = _table(doc, ["Правило", "Категория", "Уровень", "Норма"], widths)
        for rule in data.rules:
            _add_row(table, [rule.title, rule.category, SEVERITY_LABEL.get(rule.severity, rule.severity),
                             rule.legal_reference or "—"], widths, size=9)
    _footer(doc)
    doc.save(target)
    return target


def _labeled(doc, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(text)


def build_protocol(data: ReportData, target: Path) -> Path:
    """Протокол разногласий: редакция заказчика -> предлагаемая редакция -> обоснование."""
    doc = _new_document(landscape=True)
    heading = doc.add_heading("ПРОТОКОЛ РАЗНОГЛАСИЙ", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    names = ", ".join(d.file_name for d in data.documents) or data.original_filename
    p = doc.add_paragraph(f"к проекту контракта (документы: {names})")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Дата: {data.generated_at.strftime('%d.%m.%Y')}")

    widths = [1.0, 3.0, 7.4, 7.4, 6.4]
    table = _table(doc, ["№", "Пункт", "Редакция заказчика", "Предлагаемая редакция участника", "Обоснование"], widths)
    number = 0
    for rd in data.documents:
        for f in rd.issues:
            if not (f.exact_quote or f.counter_proposal):
                continue
            number += 1
            justification = f.comment + (f" ({f.legal_reference})" if f.legal_reference else "")
            location = _location(f) + (f"\n{rd.file_name}" if len(data.documents) > 1 else "")
            _add_row(table, [
                str(number), location, f.exact_quote or "Условие отсутствует",
                f.counter_proposal or "Просим уточнить / исключить условие", justification,
            ], widths, size=9)
    if number == 0:
        doc.add_paragraph("Разногласий по проверенным правилам не выявлено.")

    doc.add_paragraph()
    sign = doc.add_table(rows=2, cols=2)
    _cell_text(sign.rows[0].cells[0], "Заказчик:", bold=True)
    _cell_text(sign.rows[0].cells[1], "Участник закупки:", bold=True)
    _cell_text(sign.rows[1].cells[0], "________________ / ________________ /\nМ.П.")
    _cell_text(sign.rows[1].cells[1], "________________ / ________________ /\nМ.П.")
    _footer(doc)
    doc.save(target)
    return target


BUILDERS = {"brief": build_brief, "detailed": build_detailed, "protocol": build_protocol}
